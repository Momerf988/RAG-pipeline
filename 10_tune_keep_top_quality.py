# NEW FILE -- small keep_top quality ablation, System A only, k in {7, 15, 20}.
#
# Why this exists: 09_tune_retrieval.py already showed recall keeps climbing well past
# the point where the pre-declared +0.05 materiality bar is cleared (k=10 clears it, but
# k=15/20 clear it by a lot more). Recall alone can't tell us whether that extra retrieved
# text actually helps the tutor's real answer, or just adds noise/dilution the generator
# has to sift through. This script closes that gap with real generation + RAGAS scoring.
#
# Why System A only (not System B / CRAG): keep_top is retrieval-stage config, decided
# BEFORE either system's generation logic runs -- it's shared infrastructure, not a
# per-system setting. Testing it via System A means this result does not depend on which
# model CRAG's evaluator ends up using (that's a separate, independent decision, deferred
# until after keep_top is settled). Whatever keep_top we land on here is what both systems
# will use in the final 150-row run -- nothing here needs to be redone when the evaluator
# model question is revisited.
#
# Why k = 7, 15, 20 (not 10, not 40/60):
#   - k=7 is the current production setting AND the V7 baseline setting -- keeping it in
#     gives a direct before/after comparison point, not just three new numbers in a vacuum.
#   - k=15 and k=20 are the practical contenders from 09_tune_retrieval.py's curve (k=10
#     was ruled out as only barely clearing the materiality bar, while 15/20 clear it by
#     a much larger margin -- this ablation is what decides if that larger margin is real
#     benefit or just more text).
#   - k=40/60 are excluded here on purpose: at ~175 tokens/chunk, k=40 is already ~7,000
#     tokens of context alone, which risks exceeding Groq's free-tier 6,000 TPM cap in a
#     SINGLE generation call. That's fine for the zero-cost span_recall diagnostic (no LLM
#     call involved), but not a safe or fair setting to actually generate answers with on
#     the current infra, so it's not a realistic keep_top candidate regardless of recall.
#
# Sample: 24 questions, 6 per R-level (R1-R4), fixed seed for reproducibility -- not the
# full 150, to keep this a small/cheap check as agreed, not a repeat of the full RAGAS run.
# The SAME 24 questions are reused across all three k values, so the comparison is paired
# (same questions, different k) rather than three separate samples averaged independently.
#
# Cost: 24 questions x 3 k-values = 72 tutor generations (Groq, free tier, well under the
# 14,400/day cap) + 72 RAGAS scoring passes across 4 metrics each (OpenAI gpt-4o-mini
# judge, same setup as run_ragas_eval.py) -- a few cents total.
#
# No production files are touched. This calls RAG_response_processor's existing methods
# directly (search_db / rerank(keep_top=k) / LLM_prompt / generate_response) with keep_top
# overridden per run -- respond_to_user() itself, and its default keep_top=7, are untouched.

import os
import time
import random
import ast
import warnings
import openpyxl
import pandas as pd
from openai import RateLimitError
from datasets import Dataset
from ragas import evaluate
from ragas.metrics import faithfulness, answer_relevancy, context_precision, context_recall
from ragas.llms import LangchainLLMWrapper
from ragas.embeddings import LangchainEmbeddingsWrapper
from langchain_openai import ChatOpenAI
from langchain_community.embeddings import HuggingFaceEmbeddings

from app_init import llm_client, db_index, embedder_model, reranker_model
from user_query_processor import UserQueryProcessor
from RAG_response_processor import LLMResponseProcessor
import config

warnings.filterwarnings("ignore")

BENCHMARK_FILE = "data/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
K_VALUES = [7, 15, 20]
SAMPLE_PER_LEVEL = 6            # 6 x 4 R-levels = 24 questions total
RANDOM_SEED = 42
SLEEP_BETWEEN_ROWS = 8          # matches generate_evaluation_dataset.py -- llama-3.1-8b-instant is TPM-bottlenecked


def load_benchmark_rows():
    wb = openpyxl.load_workbook(BENCHMARK_FILE, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["Generation_Status"]] != "Generated":
            continue
        rows.append({
            "spec_id": r[idx["Spec_ID"]],
            "topic_id": r[idx["Topic_ID"]],
            "question": r[idx["Generated_Question"]],
            "ground_truth": r[idx["Reference_Answer"]],
            "r_level": r[idx["Retrieval_Level"]],
        })
    return rows


def stratified_sample(rows):
    """Same 24 questions used for every k -- paired comparison, not independent samples."""
    random.seed(RANDOM_SEED)
    by_level = {}
    for r in rows:
        by_level.setdefault(r["r_level"], []).append(r)
    sample = []
    for lvl in sorted(by_level):
        items = by_level[lvl]
        chosen = random.sample(items, min(SAMPLE_PER_LEVEL, len(items)))
        sample.extend(chosen)
    return sample


def generate_at_k(sample, k, user_query_processor, llm_response_processor):
    output_file = f"keep_top_ablation_k{k}.csv"
    if os.path.exists(output_file):
        done_df = pd.read_csv(output_file)
        done_spec_ids = set(done_df["spec_id"].tolist())
        print(f"  Resuming k={k} -- {len(done_spec_ids)} rows already done.")
    else:
        done_df = pd.DataFrame(columns=["spec_id", "topic_id", "r_level", "question",
                                         "contexts", "answer", "ground_truth"])
        done_spec_ids = set()
        print(f"  Starting fresh for k={k}.")

    for row in sample:
        spec_id = row["spec_id"]
        if spec_id in done_spec_ids:
            continue
        question = row["question"]

        retries = 0
        while retries < 3:
            try:
                embedded_query = user_query_processor.vectorize_query(question)
                db_search_result = llm_response_processor.search_db(embedded_query)
                reranked_matches = llm_response_processor.rerank(question, db_search_result, keep_top=k)
                stiched_context, context_list, chunk_id_list = llm_response_processor.stich_context(reranked_matches)
                if not stiched_context:
                    final_answer = "Sorry, there are no relevant documents in the Database to answer your query. :("
                else:
                    system_persona, user_instruction = llm_response_processor.LLM_prompt(
                        question, "elaborate", stiched_context)
                    final_answer, finish_reason = llm_response_processor.generate_response(
                        system_persona, user_instruction)
                    assert finish_reason == "stop", f"HARD ABORT: finish_reason={finish_reason}"
                break
            except RateLimitError:
                wait_seconds = 30 * (retries + 1)
                print(f"    {spec_id}: rate limit, sleeping {wait_seconds}s...")
                time.sleep(wait_seconds)
                retries += 1
        else:
            print(f"    {spec_id}: giving up after 3 retries -- rerun script to continue.")
            continue

        new_row = pd.DataFrame([{
            "spec_id": spec_id,
            "topic_id": row["topic_id"],
            "r_level": row["r_level"],
            "question": question,
            "contexts": context_list,
            "answer": final_answer,
            "ground_truth": row["ground_truth"],
        }])
        done_df = pd.concat([done_df, new_row], ignore_index=True)
        done_df.to_csv(output_file, index=False)
        done_spec_ids.add(spec_id)
        print(f"    {spec_id} done ({len(done_spec_ids)}/{len(sample)})")
        time.sleep(SLEEP_BETWEEN_ROWS)

    return output_file


def score_with_ragas(input_file, k):
    output_file = f"keep_top_ablation_k{k}_scorecard.csv"
    df = pd.read_csv(input_file)
    df["contexts"] = df["contexts"].apply(ast.literal_eval)
    df = df.rename(columns={
        "question": "user_input",
        "answer": "response",
        "contexts": "retrieved_contexts",
        "ground_truth": "reference",
    })

    judge_llm = LangchainLLMWrapper(ChatOpenAI(
        api_key=config.OPENAI_API_KEY, model="gpt-4o-mini", temperature=0, max_retries=3
    ))
    judge_embeddings = LangchainEmbeddingsWrapper(HuggingFaceEmbeddings(model_name=config.TRANSFORMER_MODEL))

    if os.path.exists(output_file):
        done_df = pd.read_csv(output_file)
        done_spec_ids = set(done_df["spec_id"].tolist())
    else:
        done_df = pd.DataFrame()
        done_spec_ids = set()

    for index in range(len(df)):
        spec_id = df.iloc[index]["spec_id"]
        if spec_id in done_spec_ids:
            continue
        single_row_df = df.iloc[[index]][["user_input", "response", "retrieved_contexts", "reference"]]
        single_eval_dataset = Dataset.from_pandas(single_row_df)
        try:
            result = evaluate(
                single_eval_dataset,
                metrics=[faithfulness, answer_relevancy, context_precision, context_recall],
                llm=judge_llm,
                embeddings=judge_embeddings,
                raise_exceptions=False,
            )
            score_dict = result.to_pandas().iloc[0].to_dict()
            score_dict["spec_id"] = spec_id
            score_dict["r_level"] = df.iloc[index]["r_level"]
            new_row = pd.DataFrame([score_dict])
            done_df = pd.concat([done_df, new_row], ignore_index=True)
            done_df.to_csv(output_file, index=False)
            done_spec_ids.add(spec_id)
        except Exception as e:
            print(f"    {spec_id}: scoring failed: {type(e).__name__}: {e}")
        time.sleep(2)

    return output_file, done_df


if __name__ == "__main__":
    print("Loading benchmark rows...")
    rows = load_benchmark_rows()
    sample = stratified_sample(rows)
    print(f"Sampled {len(sample)} questions ({SAMPLE_PER_LEVEL} per R-level), seed={RANDOM_SEED}.")
    print("Same 24 questions reused across all three k values (paired comparison, not independent samples).\n")

    user_query_processor = UserQueryProcessor(embedder_model)
    llm_response_processor = LLMResponseProcessor(llm_client, db_index, reranker_model)

    summaries = {}
    for k in K_VALUES:
        print(f"\n=== Generating answers at keep_top={k} ===")
        gen_file = generate_at_k(sample, k, user_query_processor, llm_response_processor)
        print(f"=== Scoring keep_top={k} with RAGAS (judge: gpt-4o-mini) ===")
        _, scored_df = score_with_ragas(gen_file, k)
        numeric_cols = ["faithfulness", "answer_relevancy", "context_precision", "context_recall"]
        available = [c for c in numeric_cols if c in scored_df.columns]
        summaries[k] = scored_df[available].mean()

    print("\n" + "=" * 70)
    print("SUMMARY -- keep_top quality ablation (n=24, paired sample, System A only)")
    print("=" * 70)
    summary_df = pd.DataFrame(summaries).T
    summary_df.index.name = "keep_top"
    print(summary_df)
    summary_df.to_csv("keep_top_ablation_summary.csv")
    print("\nSaved to keep_top_ablation_summary.csv")
    print("\nNOTE: this settles keep_top for BOTH System A and System B (shared retrieval")
    print("config). It does not test how keep_top affects CRAG's evaluator verdicts -- that")
    print("interaction is a disclosed scope simplification, not silently assumed away.")
