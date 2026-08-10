import pandas as pd
import time
import os
import openpyxl
from openai import RateLimitError
from app_init import llm_client, db_index, embedder_model, reranker_model
from user_query_processor import UserQueryProcessor
from RAG_response_processor import LLMResponseProcessor
from crag_evaluator import CRAGEvaluator

BENCHMARK_FILE = "data/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
OUTPUT_FILE = "System_B_eval_results.csv"
SLEEP_BETWEEN_ROWS = 12    # System B makes 2-4 LLM calls per row; give TPM more room

query_processor = UserQueryProcessor(embedder_model)
response_processor = LLMResponseProcessor(llm_client, db_index, reranker_model)
crag_evaluator = CRAGEvaluator(llm_client)

FALLBACK_MSG = "I'm sorry, that concept is outside the scope of our current syllabus or I couldn't find reliable information in the database to answer you accurately."

# Resume: load already-processed spec_ids if the file exists
if os.path.exists(OUTPUT_FILE):
    done_df = pd.read_csv(OUTPUT_FILE)
    done_spec_ids = set(done_df['spec_id'].tolist())
    print(f"Resuming — {len(done_spec_ids)} rows already done.")
else:
    done_df = pd.DataFrame(columns=[
        'spec_id', 'topic_id', 'question', 'contexts', 'answer', 'ground_truth',
        'crag_decision_1', 'rewritten_query', 'crag_decision_2', 'path_taken'
    ])
    done_spec_ids = set()
    print("Starting fresh.")

workbook = openpyxl.load_workbook(BENCHMARK_FILE)
sheet = workbook[SHEET_NAME]

def process_one_row(question):
    """Run System B's routing logic for one question.
    Returns (answer, context_list, crag_decision_1, rewritten_query, crag_decision_2, path_taken)."""
    # 1. Standard retrieval
    embedded_query = query_processor.vectorize_query(question)
    matches = response_processor.search_db(embedded_query)
    reranked_matches = response_processor.rerank(question, matches)
    context_string, context_list = response_processor.stich_context(reranked_matches)

    # 2. First CRAG eval
    decision_1 = crag_evaluator.evaluate_context(question, context_string)

    if decision_1 == "CORRECT":
        # Direct answer path
        system_persona, user_instruction = response_processor.LLM_prompt(question, 'elaborate', context_string)
        answer = response_processor.generate_response(system_persona, user_instruction)
        return answer, context_list, decision_1, "", "", "direct"

    # 3. Corrective path: rewrite + second retrieval
    rewritten = crag_evaluator.rewrite_query(question)
    embedded_rewritten = query_processor.vectorize_query(rewritten)
    matches_2 = response_processor.search_db(embedded_rewritten)
    reranked_matches_2 = response_processor.rerank(rewritten, matches_2)
    context_string_2, context_list_2 = response_processor.stich_context(reranked_matches_2)

    # 4. Second CRAG eval
    decision_2 = crag_evaluator.evaluate_context(rewritten, context_string_2)

    if decision_2 == "CORRECT":
        system_persona, user_instruction = response_processor.LLM_prompt(question, 'elaborate', context_string_2)
        answer = response_processor.generate_response(system_persona, user_instruction)
        return answer, context_list_2, decision_1, rewritten, decision_2, "rewritten"

    # 5. Graceful fallback
    return FALLBACK_MSG, context_list_2, decision_1, rewritten, decision_2, "fallback"


print("Starting System B evaluation loop...")

for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
    spec_id, topic_id = row[0], row[1]
    status = row[9]
    question = row[10]
    ground_truth = row[11]

    if status != "Generated" or not question:
        continue
    if spec_id in done_spec_ids:
        continue

    # rate-limit-safe call with retry
    retries = 0
    result = None
    while retries < 3:
        try:
            result = process_one_row(question)
            break
        except RateLimitError as e:
            wait_seconds = 30 * (retries + 1)
            print(f"{spec_id}: rate limit hit, sleeping {wait_seconds}s...")
            time.sleep(wait_seconds)
            retries += 1
    if result is None:
        print(f"{spec_id}: giving up after 3 retries. Progress saved; rerun to continue.")
        break

    answer, context_list, decision_1, rewritten, decision_2, path = result

    new_row = pd.DataFrame([{
        'spec_id': spec_id,
        'topic_id': topic_id,
        'question': question,
        'contexts': context_list,
        'answer': answer,
        'ground_truth': ground_truth,
        'crag_decision_1': decision_1,
        'rewritten_query': rewritten,
        'crag_decision_2': decision_2,
        'path_taken': path
    }])
    done_df = pd.concat([done_df, new_row], ignore_index=True)
    done_df.to_csv(OUTPUT_FILE, index=False)

    done_spec_ids.add(spec_id)
    print(f"Processed {spec_id} — path: {path}, decision_1: {decision_1}. Total done: {len(done_spec_ids)}")
    time.sleep(SLEEP_BETWEEN_ROWS)

print(f"\nDone. {len(done_spec_ids)} rows saved to {OUTPUT_FILE}.")

# Path breakdown for quick sanity
if len(done_df) > 0:
    print("\nPath breakdown:")
    print(done_df['path_taken'].value_counts())