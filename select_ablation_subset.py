import openpyxl
import random

BENCHMARK_FILE = "data/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
OUTPUT_FILE = "ablation_subset_spec_ids.txt"
TARGET_SIZE = 25

random.seed(42)  # reproducible sample

workbook = openpyxl.load_workbook(BENCHMARK_FILE)
sheet = workbook[SHEET_NAME]

rows_by_topic = {}
for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
    spec_id, topic_id, status = row[0], row[1], row[9]
    if status != "Generated":
        continue
    rows_by_topic.setdefault(topic_id, []).append(spec_id)

topics = sorted(rows_by_topic.keys())
per_topic = max(1, TARGET_SIZE // len(topics))

selected = []
for topic in topics:
    pool = rows_by_topic[topic]
    n = min(per_topic, len(pool))
    selected.extend(random.sample(pool, n))

# top up to TARGET_SIZE if short, from remaining unselected rows
remaining = [sid for topic in topics for sid in rows_by_topic[topic] if sid not in selected]
while len(selected) < TARGET_SIZE and remaining:
    pick = random.choice(remaining)
    selected.append(pick)
    remaining.remove(pick)

selected = selected[:TARGET_SIZE]

with open(OUTPUT_FILE, "w") as f:
    for sid in selected:
        f.write(sid + "\n")

print(f"Selected {len(selected)} spec_ids for ablation, saved to {OUTPUT_FILE}")
print(f"Breakdown by topic:")
for topic in topics:
    count = sum(1 for sid in selected if sid.startswith(topic + "_"))
    print(f"  {topic}: {count}")