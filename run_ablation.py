import pandas as pd
import time
import openpyxl
from openai import RateLimitError
from app_init import llm_client, db_index, embedder_model, reranker_model
from user_query_processor import UserQueryProcessor
from RAG_response_processor import LLMResponseProcessor

# --- EDIT THIS BEFORE EACH OF THE 3 RUNS ---
KEEP_TOP = 7   # run once each with 3, 5, 7

BENCHMARK_FILE = "data/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
SUBSET_FILE = "ablation_subset_spec_ids.txt"
OUTPUT_FILE = f"ablation_keep_top_{KEEP_TOP}.csv"

with open(SUBSET_FILE) as f:
    subset_ids = set(line.strip() for line in f if line.strip())

query_processor = UserQueryProcessor(embedder_model)
response_processor = LLMResponseProcessor(llm_client, db_index, reranker_model)

workbook = openpyxl.load_workbook(BENCHMARK_FILE)
sheet = workbook[SHEET_NAME]

eval_data = []
print(f"Running ablation with keep_top={KEEP_TOP} on {len(subset_ids)} rows...")

for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
    spec_id, topic_id = row[0], row[1]
    status = row[9]
    question = row[10]
    ground_truth = row[11]

    if spec_id not in subset_ids:
        continue
    if status != "Generated" or not question:
        continue

    retries = 0
    while retries < 3:
        try:
            embedded_query = query_processor.vectorize_query(question)
            matches = response_processor.search_db(embedded_query)
            reranked = response_processor.rerank(question, matches, keep_top=KEEP_TOP)
            context_string, context_list = response_processor.stich_context(reranked)

            if not context_string:
                final_answer = "Sorry, there are no relevant documents in the Database to answer your query. :("
            else:
                system_persona, user_instruction = response_processor.LLM_prompt(question, 'elaborate', context_string)
                final_answer = response_processor.generate_response(system_persona, user_instruction)
            break
        except RateLimitError:
            wait_seconds = 30 * (retries + 1)
            print(f"{spec_id}: rate limit, sleeping {wait_seconds}s...")
            time.sleep(wait_seconds)
            retries += 1
    else:
        print(f"{spec_id}: giving up after 3 retries.")
        continue

    eval_data.append({
        'spec_id': spec_id,
        'topic_id': topic_id,
        'question': question,
        'contexts': context_list,
        'answer': final_answer,
        'ground_truth': ground_truth
    })
    print(f"Processed {spec_id}")
    time.sleep(8)

eval_results_df = pd.DataFrame(eval_data)
eval_results_df.to_csv(OUTPUT_FILE, index=False)
print(f"\nDone. Saved {len(eval_data)} rows to {OUTPUT_FILE}")