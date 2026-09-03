import json
import random
import time
import openpyxl
from app_init import llm_client, db_index
import config

# Generates the benchmark questions, one topic at a time, from the spec matrix's locked
# rows (Bloom level, difficulty, retrieval level, question type, programming context all
# pre-assigned before any question exists). Run once per topic, edit TOPIC_TO_GENERATE
# below and rerun for the next one -- keeping it manual per topic makes it easy to spot-
# check a topic's output before moving on to the next.

TOPIC_TO_GENERATE = "T3"

SPEC_FILE = "benchmark_datasets/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
GENERATOR_MODEL = "openai/gpt-oss-120b"

RETRIEVAL_CHUNK_COUNTS = {
    "R1": 1,
    "R2": 2,
    "R3": 3,
    "R4": 4
}


class BenchmarkQuestionGenerator:
    def __init__(self, llm_client, db_index):
        self.llm_client = llm_client
        self.index = db_index

    def get_topic_chunk_pool(self, source_pdf):
        # Vector IDs were written as f"{pdf_file}_{i}" during ingestion, so the PDF name is
        # a usable prefix -- no need to touch the ingestion schema to look chunks up by topic.
        prefix = f"{source_pdf}_"
        chunk_pool = {}
        for id_batch in self.index.list(prefix=prefix):
            fetched = self.index.fetch(ids=id_batch)
            for vector_id, vector_data in fetched.vectors.items():
                chunk_pool[vector_id] = vector_data.metadata['line']
        return chunk_pool

    def pick_chunks(self, chunk_pool, retrieval_level):
        # Sort by the numeric suffix in the vector ID so "sequential" actually means
        # sequential in the source document, then take a contiguous run of the length the
        # retrieval level calls for, starting at a random position.
        chunk_ids = sorted(chunk_pool.keys(), key=lambda cid: int(cid.rsplit("_", 1)[-1]))

        chunk_count = RETRIEVAL_CHUNK_COUNTS.get(retrieval_level, 1)
        chunk_count = min(chunk_count, len(chunk_ids))

        max_start_idx = max(0, len(chunk_ids) - chunk_count)
        start_idx = random.randint(0, max_start_idx)

        sampled_ids = chunk_ids[start_idx: start_idx + chunk_count]
        sampled_text = [chunk_pool[cid] for cid in sampled_ids]

        return sampled_text, sampled_ids

    def build_generation_prompt(self, topic, bloom, difficulty, question_type, programming_context, context_string):
        system_persona = "You are an expert university-level Python instructor writing exam questions for an Intelligent Tutoring System."
        user_instruction = f'''
Write ONE exam question that tests the spec below, grounded ONLY in the Context provided.

Topic: {topic}
Bloom's Level: {bloom}
Difficulty: {difficulty}
Question_Type: {question_type}
Programming_Context: {programming_context}

CRITICAL RULES:
- Match the Bloom level exactly (e.g. "Analyse" must ask the student to break down/compare/examine, not just recall a fact).
- Match Question_Type exactly (e.g. "Debugging" = show a broken snippet and ask to find/fix the bug; "Code Tracing" = give code and ask what it outputs; "Definition" = ask what a term/concept means).
- If Programming_Context is "Code" or "Mixed", the question must include an actual Python code snippet.
- THE "PHANTOM SNIPPET" RULE: If asking a "Comparison" or "Debugging" question, you MUST write the actual Python code snippets directly inside the question string. Do not say "the following code" without actually providing the code.
- THE "REAL BUG" RULE: If the type is "Debugging", take the correct code shown in the Context and rewrite it with ONE deliberate, realistic bug you introduce yourself (off-by-one, wrong operator, wrong variable). State clearly in the question that the code is broken. The Reference_Answer must name the exact line/operator you changed and show the corrected version.
- THE "EXECUTION" RULE: Double-check all string slicing, indexing, and math in your head. For example, word[1:3] on 'Python' is 'yt', not 'th'. Your reference answer must be 100% factually accurate.
- THE "SCOPED EXAMPLE" RULE: If constructing a scenario that involves a loop, sequence, or multi-step calculation, choose bounds (e.g. "less than 20" instead of "less than 1000") so that any full trace needs at most 5-6 steps. Never construct a scenario that requires 10+ iterations to answer correctly.
- THE "CONCISE TRUTH" RULE: Keep the Reference_Answer to 1-3 short paragraphs by default. The ONLY exception is Code Tracing or Debugging questions, where step-by-step work (a short table or numbered steps) is allowed because it's often necessary to prove correctness -- but even then, keep it within the scope allowed by the SCOPED EXAMPLE rule above.
- The Reference_Answer must be fully answerable using ONLY the Context below -- do not invent facts that aren't in it.
- Reply with ONLY a JSON object, no markdown fences, no extra text, in exactly this shape:

{{"question": "...", "answer": "...", "section": "short phrase naming the concept this targets"}}

Context:
{context_string}
'''
        # Only fires for the Understand x R3/R4 rows, which use a question type the other
        # rows don't: forces the answer to draw on every supplied chunk instead of just one,
        # so the row actually tests wide evidence rather than a single lucky chunk.
        if question_type == "Enumeration":
            user_instruction += '''

ENUMERATION RULE: The question must ask the student to list or enumerate information explicitly stated across the supplied context.
The answer must require information from the supplied R-level context, not from outside knowledge or from only a single unrelated portion of the context.
Keep the cognitive demand at recall/comprehension only -- do not ask the student to judge, compare, evaluate, or analyse anything.'''

        return system_persona, user_instruction

    def generate_response(self, system_persona, user_instruction):
        response = self.llm_client.chat.completions.create(
            model=GENERATOR_MODEL,
            messages=[
                {"role": "system", "content": system_persona},
                {"role": "user", "content": user_instruction}
            ],
            temperature=0.3,
            max_tokens=4000,  # long R4 Evaluate answers need the headroom
            response_format={"type": "json_object"},
        )
        return response.choices[0].message.content

    def parse_generated_json(self, raw_text):
        # Slicing to the outer braces (rather than trusting the model to skip markdown
        # fences cleanly) handles the occasional response that wraps the JSON in commentary
        # or code fences despite being told not to.
        start = raw_text.find("{")
        end = raw_text.rfind("}")
        if start == -1 or end == -1:
            raise ValueError("No JSON object found in model response")
        cleaned = raw_text[start:end + 1]
        parsed = json.loads(cleaned, strict=False)
        return parsed["question"], parsed["answer"], parsed["section"]

    def generate_for_topic(self, topic_id):
        workbook = openpyxl.load_workbook(SPEC_FILE)
        sheet = workbook[SHEET_NAME]

        COL_SPEC_ID, COL_TOPIC_ID, COL_TOPIC, COL_SOURCE_PDF = 1, 2, 3, 4
        COL_BLOOM, COL_DIFFICULTY, COL_RETRIEVAL, COL_QTYPE, COL_PCTX = 5, 6, 7, 8, 9
        COL_STATUS, COL_QUESTION, COL_ANSWER, COL_CHUNKS, COL_SECTION = 10, 11, 12, 13, 14

        chunk_pool = None
        generated_count = 0
        skipped_count = 0

        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            if row[COL_TOPIC_ID - 1].value != topic_id:
                continue
            if row[COL_STATUS - 1].value != "Pending":
                continue

            spec_id = row[COL_SPEC_ID - 1].value
            source_pdf = row[COL_SOURCE_PDF - 1].value

            if chunk_pool is None:
                print(f"Fetching chunk pool for {source_pdf} ...")
                chunk_pool = self.get_topic_chunk_pool(source_pdf)
                print(f"Found {len(chunk_pool)} chunks for {topic_id}.")

            retrieval_level = row[COL_RETRIEVAL - 1].value
            context_texts, context_ids = self.pick_chunks(chunk_pool, retrieval_level)
            context_string = " ".join(context_texts)

            system_persona, user_instruction = self.build_generation_prompt(
                row[COL_TOPIC - 1].value,
                row[COL_BLOOM - 1].value,
                row[COL_DIFFICULTY - 1].value,
                row[COL_QTYPE - 1].value,
                row[COL_PCTX - 1].value,
                context_string
            )

            try:
                raw_response = self.generate_response(system_persona, user_instruction)
                question, answer, section = self.parse_generated_json(raw_response)

                row[COL_QUESTION - 1].value = question
                row[COL_ANSWER - 1].value = answer
                row[COL_CHUNKS - 1].value = ", ".join(context_ids)
                row[COL_SECTION - 1].value = section
                row[COL_STATUS - 1].value = "Generated"

                generated_count += 1
                print(f"{spec_id}: generated.")
            except Exception as e:
                skipped_count += 1
                print(f"{spec_id}: skipped -- {type(e).__name__}: {e}")

            time.sleep(2)  # respect Groq free-tier rate limits

        workbook.save(SPEC_FILE)
        print(f"\nTopic {topic_id}: {generated_count} generated, {skipped_count} skipped.")


if __name__ == "__main__":
    generator = BenchmarkQuestionGenerator(llm_client, db_index)
    generator.generate_for_topic(TOPIC_TO_GENERATE)
