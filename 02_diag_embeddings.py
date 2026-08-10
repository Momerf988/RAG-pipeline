# NEW FILE -- embedding model ablation (REBUILD_SPEC Step 02, adapted to your actual scope).
#
# Purpose: answer one question -- "is the current embedder (MiniLM) capping retrieval
# quality, or would a stronger embedder not actually help?" -- with a measured number
# instead of a guess. This is the comparison you agreed to run after the list/tuple case.
#
# What it does NOT do (deliberately, matching what you've already scoped):
#   - No LLM calls, no cost. Pure local encoding + cosine similarity + your existing
#     cross-encoder reranker.
#   - No Pinecone writes, no index changes. Fully offline, using numpy.
#   - No dev/test split (that was explicitly deferred). Uses the full 150-row benchmark
#     directly via the gold_spans.json you already built with spanify_gold.py.
#   - Does not touch chunking. Same 317 chunks, same chunk() function, verbatim, as
#     document_processor.py. Only the embedding model changes between the two runs.
#
# Decision rule, declared BEFORE looking at the numbers (REBUILD_SPEC Step 02):
#   switch to bge-base-en-v1.5 only if overall recall improves by >= +0.05,
#   OR R4 recall improves by >= +0.10. Otherwise keep MiniLM -- and that's a legitimate,
#   reportable result, not a failed experiment.

import os
import json
import openpyxl
import numpy as np
from sentence_transformers import SentenceTransformer, CrossEncoder
import config

PDF_DIR = "pdf_files"
SPEC_FILE = "data/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
GOLD_SPANS_FILE = "data/gold_spans.json"

CANDIDATES = [
    config.TRANSFORMER_MODEL,     # incumbent -- whatever is actually configured, not assumed
    "BAAI/bge-base-en-v1.5",      # the honest upgrade candidate named in REBUILD_SPEC
]
RERANKER_MODEL = "cross-encoder/ms-marco-MiniLM-L-6-v2"  # unchanged, same as production
TOP_K_CANDIDATES = 20   # pre-rerank pool size, matches RAG_response_processor.search_db
KEEP_TOP = 7            # post-rerank, matches RAG_response_processor.rerank


# ---- chunking: copied verbatim from document_processor.py. Do not edit. ----
def read_pdf(file_path):
    import pdfplumber
    raw_text = ""
    with pdfplumber.open(file_path) as pdf_file:
        for page in pdf_file.pages:
            extracted_text = page.extract_text()
            if extracted_text:
                raw_text += extracted_text + "\n"
    return raw_text


def chunking(raw_text, chunk_size=700, overlap=100):
    raw_text = raw_text.replace("(cid:415)", "ti")
    chunks = []
    start = 0
    text_length = len(raw_text)
    while start < text_length:
        end = start + chunk_size
        if end < text_length:
            newline_pos = raw_text.rfind("\n", start, end)
            if newline_pos > start:
                end = newline_pos
            else:
                sentence_pos = raw_text.rfind(". ", start, end)
                if sentence_pos > start:
                    end = sentence_pos + 1
                else:
                    space_pos = raw_text.rfind(" ", start, end)
                    if space_pos > start:
                        end = space_pos
        chunk_group = raw_text[start:end].strip("\n").strip()
        if chunk_group:
            chunks.append(chunk_group)
        next_start = end - overlap if end - overlap > start else end
        if next_start < text_length:
            space_pos = raw_text.find(" ", next_start)
            newline_pos = raw_text.find("\n", next_start)
            candidates = [p for p in (space_pos, newline_pos) if p != -1]
            if candidates:
                boundary = min(candidates)
                if boundary - next_start < 50:
                    next_start = boundary + 1
        start = next_start
    return chunks
# ---- end verbatim copy ----


def build_corpus():
    """Rebuild the 317 chunks with their (start, end) character offsets per source file."""
    pdf_files = sorted(f for f in os.listdir(PDF_DIR) if f.endswith(".pdf"))
    all_chunk_ids, all_chunk_texts = [], []
    offsets_by_chunk_id = {}
    total = 0
    for pdf_file in pdf_files:
        raw_text = read_pdf(os.path.join(PDF_DIR, pdf_file))
        chunks = chunking(raw_text, chunk_size=700, overlap=100)
        cursor = 0
        for i, chunk_text in enumerate(chunks):
            cid = f"{pdf_file}_{i}"
            start = raw_text.index(chunk_text, cursor)
            end = start + len(chunk_text)
            cursor = start + 1
            offsets_by_chunk_id[cid] = (start, end)
            all_chunk_ids.append(cid)
            all_chunk_texts.append(chunk_text)
        total += len(chunks)
    assert total == config.EXPECTED_CHUNKS, \
        f"HARD ABORT: rebuilt {total} chunks, expected {config.EXPECTED_CHUNKS}"
    print(f"Corpus OK: {total} chunks across {len(pdf_files)} files.")
    return all_chunk_ids, all_chunk_texts, offsets_by_chunk_id


def load_benchmark_rows():
    """spec_id -> (question, retrieval_level, source_pdf)"""
    wb = openpyxl.load_workbook(SPEC_FILE, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["Generation_Status"]] != "Generated":
            continue
        rows[r[idx["Spec_ID"]]] = {
            "question": r[idx["Generated_Question"]],
            "r_level": r[idx["Retrieval_Level"]],
            "source_pdf": r[idx["Source_PDF"]],
        }
    return rows


def merge_intervals(intervals):
    if not intervals:
        return []
    s = sorted(intervals)
    merged = [list(s[0])]
    for cur_s, cur_e in s[1:]:
        if cur_s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], cur_e)
        else:
            merged.append([cur_s, cur_e])
    return [tuple(m) for m in merged]


def overlap_len(a, b):
    return max(0, min(a[1], b[1]) - max(a[0], b[0]))


def span_recall(gold_spans_list, retrieved_spans_list):
    """Fraction of gold-span characters covered by the retrieved chunks' character ranges."""
    total_gold_chars = sum(e - s for s, e in gold_spans_list)
    if total_gold_chars == 0:
        return None
    merged_retrieved = merge_intervals(retrieved_spans_list)
    covered = 0
    for g in gold_spans_list:
        for r in merged_retrieved:
            covered += overlap_len(g, r)
    return covered / total_gold_chars


def run_model(model_name, chunk_ids, chunk_texts, offsets_by_chunk_id, bench_rows,
              gold_spans, reranker):
    print(f"\nEncoding {len(chunk_texts)} chunks with {model_name} ...")
    embedder = SentenceTransformer(model_name)
    chunk_vecs = embedder.encode(chunk_texts, normalize_embeddings=True, show_progress_bar=False)
    chunk_vecs = np.array(chunk_vecs)

    results = []  # (spec_id, r_level, recall)
    for spec_id, gold in gold_spans.items():
        if spec_id not in bench_rows:
            continue
        question = bench_rows[spec_id]["question"]
        r_level = bench_rows[spec_id]["r_level"]
        source_pdf = gold["source"]
        gold_spans_list = [(s["start"], s["end"]) for s in gold["spans"]]

        q_vec = embedder.encode([question], normalize_embeddings=True)[0]
        sims = chunk_vecs @ q_vec
        top20_idx = np.argsort(-sims)[:TOP_K_CANDIDATES]
        top20_ids = [chunk_ids[i] for i in top20_idx]
        top20_texts = [chunk_texts[i] for i in top20_idx]

        pairs = [(question, t) for t in top20_texts]
        scores = reranker.predict(pairs)
        reranked = sorted(zip(scores, top20_ids), key=lambda x: x[0], reverse=True)[:KEEP_TOP]
        top7_ids = [cid for _, cid in reranked]

        # only chunks from the SAME source file are meaningful for overlap -- offsets are
        # file-local, comparing across files would be comparing unrelated coordinate systems
        retrieved_spans = [offsets_by_chunk_id[cid] for cid in top7_ids if cid.startswith(source_pdf)]

        recall = span_recall(gold_spans_list, retrieved_spans)
        if recall is not None:
            results.append((spec_id, r_level, recall))

    return results


def summarize(model_name, results):
    overall = np.mean([r for _, _, r in results])
    by_level = {}
    for _, lvl, r in results:
        by_level.setdefault(lvl, []).append(r)
    print(f"\n=== {model_name} ===")
    print(f"  Overall span-recall@{KEEP_TOP}: {overall:.4f}  (n={len(results)})")
    for lvl in sorted(by_level):
        vals = by_level[lvl]
        print(f"  {lvl}: {np.mean(vals):.4f}  (n={len(vals)})")
    return overall, {lvl: np.mean(v) for lvl, v in by_level.items()}


if __name__ == "__main__":
    print("Loading benchmark and gold spans...")
    bench_rows = load_benchmark_rows()
    with open(GOLD_SPANS_FILE) as f:
        gold_spans = json.load(f)
    print(f"{len(bench_rows)} generated rows, {len(gold_spans)} gold spans loaded.")

    chunk_ids, chunk_texts, offsets_by_chunk_id = build_corpus()

    print(f"\nLoading reranker ({RERANKER_MODEL}) -- shared, unchanged across both models...")
    reranker = CrossEncoder(RERANKER_MODEL)

    all_summaries = {}
    for model_name in CANDIDATES:
        results = run_model(model_name, chunk_ids, chunk_texts, offsets_by_chunk_id,
                             bench_rows, gold_spans, reranker)
        overall, by_level = summarize(model_name, results)
        all_summaries[model_name] = {"overall": overall, "by_level": by_level}

    # ---- apply the pre-declared decision rule ----
    incumbent, candidate = CANDIDATES[0], CANDIDATES[1]
    inc_overall = all_summaries[incumbent]["overall"]
    cand_overall = all_summaries[candidate]["overall"]
    inc_r4 = all_summaries[incumbent]["by_level"].get("R4")
    cand_r4 = all_summaries[candidate]["by_level"].get("R4")

    overall_gain = cand_overall - inc_overall
    r4_gain = (cand_r4 - inc_r4) if (inc_r4 is not None and cand_r4 is not None) else None

    print("\n" + "=" * 60)
    print("DECISION (pre-declared rule, REBUILD_SPEC Step 02):")
    print(f"  Overall gain: {overall_gain:+.4f}  (switch threshold: >= +0.05)")
    if r4_gain is not None:
        print(f"  R4 gain:      {r4_gain:+.4f}  (switch threshold: >= +0.10)")
    if overall_gain >= 0.05 or (r4_gain is not None and r4_gain >= 0.10):
        print(f"  -> SWITCH to {candidate}. Threshold met.")
    else:
        print(f"  -> KEEP {incumbent}. Threshold not met -- this is a valid, reportable outcome.")
    print("=" * 60)
