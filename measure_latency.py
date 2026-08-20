# V2.6: fixed before the overnight run -- this whole script was stale relative to several
# already-completed changes and would have crashed on the first call:
#   1. stich_context() now returns 3 values (added chunk_id_list); this script still
#      unpacked 2, everywhere it called it -- "ValueError: too many values to unpack".
#   2. generate_response() now returns (answer, finish_reason); one call site here captured
#      the tuple into a single "final_answer" var (unused afterward so not a crash, but
#      wrong), cleaned up for correctness.
#   3. OUTPUT_FILE used f"...{{T}}..." (double braces) inside an f-string, which produces a
#      literal "{T}" in the filename instead of interpolating the variable -- fixed to single
#      braces.
#   4. System B's generation calls used LLM_prompt (System A's prompt), not LLM_prompt_crag
#      (System B's actual production prompt, built later than this script). Fixed so the
#      timed call matches what System B really runs -- otherwise this would time a call B
#      never actually makes.
#
# V2.11: this script was the one piece of the pipeline with no resume/retry safety net --
# confirmed live when it hit the account's daily token cap (429, tokens-per-day) and crashed
# uncaught, losing every row it had already timed since results were only written to disk
# once, at the very end of main(). Now matches the same pattern already used in
# generate_evaluation_dataset.py / _B.py: reads its own OUTPUT_FILE on startup and skips
# spec_ids already timed, saves after every row (not just at the end), retries RateLimitError
# 3x with backoff then stops cleanly (progress kept) instead of crashing, and skips
# (not crashes on) AssertionError/APIStatusError the same way the generation scripts do. This
# is what makes it safe to run as its own separate step -- interrupt it any time (deadline,
# quota, closing the laptop) and rerunning just continues from wherever it left off.
import time
import random
import os
import openpyxl
import pandas as pd
from openai import RateLimitError, APIStatusError
from app_init import llm_client, db_index, embedder_model, reranker_model
from user_query_processor import UserQueryProcessor
from RAG_response_processor import LLMResponseProcessor
from crag_evaluator import CRAGEvaluator
# V2.14: T (inter-row pacing delay) was run at 1, 30 and 70s. This is NOT an architecture
# parameter -- it only controls request spacing against OpenRouter, and all three runs
# target the identical 20-row sample (random.seed(42) below). T=1 showed heavy contention
# (single evaluator calls up to 33s; 4-5 LLM calls fired per row with almost no gap), which
# eased substantially by T=70 -- consistent with provider-side queueing/throttling, not a
# change in what either system computes. T=70 is used as the primary reported run since it
# has the least contention, but even it isn't fully clean (still real outliers), so report
# MEDIAN alongside mean when writing this up -- the median overhead (~0.9s) lines up closely
# with V7's own ~1s finding, while the mean is pulled up by remaining API-side variance.
# Do not describe T=70 as "clearing the rate limit" -- it only reduces the likelihood of
# contention, it doesn't guarantee it.
T = 70
BENCHMARK_FILE = "data/benchmark_specification_matrix.xlsx"
SHEET_NAME = "Benchmark Spec Matrix"
OUTPUT_FILE = f"Latency_Results_t{T}.csv"
SAMPLE_SIZE_PER_RETRIEVAL_LEVEL = 5   # 5 rows x 4 levels (R1-R4) = 20 rows total

random.seed(42)

query_processor = UserQueryProcessor(embedder_model)
response_processor = LLMResponseProcessor(llm_client, db_index, reranker_model)
crag_evaluator = CRAGEvaluator(llm_client)


def select_latency_sample():
    workbook = openpyxl.load_workbook(BENCHMARK_FILE)
    sheet = workbook[SHEET_NAME]

    rows_by_retrieval = {}
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
        spec_id, topic_id, status, question, retrieval_level = row[0], row[1], row[9], row[10], row[6]
        if status != "Generated" or not question:
            continue
        rows_by_retrieval.setdefault(retrieval_level, []).append((spec_id, topic_id, question))

    sample = []
    for level, rows in rows_by_retrieval.items():
        n = min(SAMPLE_SIZE_PER_RETRIEVAL_LEVEL, len(rows))
        sample.extend([(*r, level) for r in random.sample(rows, n)])
    return sample


def time_system_a(question):
    t0 = time.perf_counter()
    embedded_query = query_processor.vectorize_query(question)
    matches = response_processor.search_db(embedded_query)
    reranked = response_processor.rerank(question, matches)
    context_string, context_list, chunk_ids = response_processor.stich_context(reranked)
    t1 = time.perf_counter()

    if not context_string:
        return {"retrieval_ms": (t1 - t0) * 1000, "evaluator_ms": 0, "generation_ms": 0, "total_ms": (t1 - t0) * 1000}

    system_persona, user_instruction = response_processor.LLM_prompt(question, 'elaborate', context_string)
    final_answer, finish_reason = response_processor.generate_response(system_persona, user_instruction)
    t2 = time.perf_counter()
    time.sleep(5)  # clear rate-limit headroom before the next LLM call anywhere in the script

    return {
        "retrieval_ms": (t1 - t0) * 1000,
        "evaluator_ms": 0,
        "generation_ms": (t2 - t1) * 1000,
        "total_ms": (t2 - t0) * 1000,
    }


# V2.12: fixed the same instrumentation fault the V7 report already documented and had to
# correct for after the fact (Section 4.7/5.3 of V7.docx: "fixed five-second pauses inserted
# between chained model calls fall inside whichever window brackets them"). It was never
# actually fixed in the code -- V7 subtracted known pause constants post-hoc instead -- so it
# silently carried forward into this rebuild. Confirmed live against Latency_Results_t1.csv:
# direct-path total_ms residual was 5005ms (std 2.9ms) and rewritten-path rewrite_cycle_ms
# carried an extra ~10,000ms -- both dead-on matches for V7's own reported 5,004ms / 10,010ms
# figures. Root cause: every time.sleep(5) here was placed to pace API calls against the
# provider's rate limit, but several of them sat BETWEEN a segment's start and end timestamp
# (e.g. generation_ms = t3-t2 while a sleep(5) ran between t2 being captured and generate_
# response() actually being called), so the pacing delay got silently counted as model work.
# System A's timing was never affected because its one sleep(5) already came after its final
# timestamp was captured -- the fix here is to make System B's do the same: capture every
# timestamp immediately around its real API call, and move ALL pacing sleeps to run only
# after every timestamp for the row has already been taken, right before returning. total_ms
# is also now the sum of the clean components rather than a fresh wall-clock read, so it can
# never re-absorb a pacing delay no matter where one is placed.
def time_system_b(question):
    t0 = time.perf_counter()
    embedded_query = query_processor.vectorize_query(question)
    matches = response_processor.search_db(embedded_query)
    reranked = response_processor.rerank(question, matches)
    context_string, context_list, chunk_ids = response_processor.stich_context(reranked)
    t1 = time.perf_counter()

    decision_1 = crag_evaluator.evaluate_context(question, context_string)
    t2 = time.perf_counter()

    retrieval_ms = (t1 - t0) * 1000
    evaluator_ms = (t2 - t1) * 1000
    generation_ms = 0
    rewrite_ms = 0
    path = "direct"

    if decision_1 == "CORRECT":
        # matches production: System B generates with LLM_prompt_crag, not LLM_prompt
        system_persona, user_instruction = response_processor.LLM_prompt_crag(question, 'elaborate', context_string)
        response_processor.generate_response(system_persona, user_instruction)
        t3 = time.perf_counter()
        generation_ms = (t3 - t2) * 1000
    else:
        path = "rewritten_or_fallback"
        rewritten = crag_evaluator.rewrite_query(question)
        t_rewrite_done = time.perf_counter()
        embedded_rewritten = query_processor.vectorize_query(rewritten)
        matches_2 = response_processor.search_db(embedded_rewritten)
        reranked_2 = response_processor.rerank(rewritten, matches_2)
        context_string_2, context_list_2, chunk_ids_2 = response_processor.stich_context(reranked_2)
        decision_2 = crag_evaluator.evaluate_context(rewritten, context_string_2)
        t3 = time.perf_counter()
        # rewrite_ms = time from right after the evaluator's first verdict to the second
        # verdict landing -- rewrite call + second retrieval + second evaluator call, no
        # sleeps anywhere inside this span now.
        rewrite_ms = (t3 - t2) * 1000

        if decision_2 in ("CORRECT", "AMBIGUOUS"):
            # matches production routing (system_b_main.py / generate_evaluation_dataset_B.py):
            # CORRECT and AMBIGUOUS both generate on the second pass, only a second INCORRECT
            # falls back. Also matches production: LLM_prompt_crag, not LLM_prompt.
            system_persona, user_instruction = response_processor.LLM_prompt_crag(question, 'elaborate', context_string_2)
            response_processor.generate_response(system_persona, user_instruction)
            t4 = time.perf_counter()
            generation_ms = (t4 - t3) * 1000
        else:
            t4 = t3

    # total_ms is derived from the already-clean components, not a fresh wall-clock read --
    # this is what actually guarantees it can't reabsorb the pacing sleep below, regardless
    # of where that sleep is placed.
    total_ms = retrieval_ms + evaluator_ms + rewrite_ms + generation_ms
    result = {
        "retrieval_ms": retrieval_ms,
        "evaluator_ms": evaluator_ms,
        "rewrite_cycle_ms": rewrite_ms,
        "generation_ms": generation_ms,
        "total_ms": total_ms,
        "path": path,
    }
    time.sleep(5)  # rate-limit pacing before the next row -- runs after every timestamp for
                    # this row has already been captured, so it can't contaminate anything above
    return result


def main():
    sample = select_latency_sample()

    # Resume: load already-timed spec_ids if the file exists (same pattern as the two
    # generation scripts) -- running this script again after any interruption continues
    # instead of re-timing rows or losing what's already saved.
    if os.path.exists(OUTPUT_FILE):
        existing_df = pd.read_csv(OUTPUT_FILE)
        done_spec_ids = set(existing_df['spec_id'].tolist())
        results = existing_df.to_dict('records')
        print(f"Resuming -- {len(done_spec_ids)} rows already timed.")
    else:
        done_spec_ids = set()
        results = []
        print("Starting fresh.")

    remaining = [row for row in sample if row[0] not in done_spec_ids]
    print(f"Timing {len(sample)} rows total across both systems ({len(remaining)} remaining)...\n")

    for spec_id, topic_id, question, retrieval_level in remaining:
        print(f"Timing {spec_id} ({retrieval_level})...")

        retries = 0
        row_result = None
        skipped = False
        while retries < 3:
            try:
                a_timing = time_system_a(question)
                b_timing = time_system_b(question)
                row_result = {
                    "spec_id": spec_id,
                    "topic_id": topic_id,
                    "retrieval_level": retrieval_level,
                    "A_retrieval_ms": a_timing["retrieval_ms"],
                    "A_generation_ms": a_timing["generation_ms"],
                    "A_total_ms": a_timing["total_ms"],
                    "B_retrieval_ms": b_timing["retrieval_ms"],
                    "B_evaluator_ms": b_timing["evaluator_ms"],
                    "B_rewrite_cycle_ms": b_timing.get("rewrite_cycle_ms", 0),
                    "B_generation_ms": b_timing["generation_ms"],
                    "B_total_ms": b_timing["total_ms"],
                    "B_path": b_timing["path"],
                }
                break
            except RateLimitError as e:
                wait_seconds = 30 * (retries + 1)   # 30s, 60s, 90s
                print(f"{spec_id}: rate limit hit, sleeping {wait_seconds}s...")
                time.sleep(wait_seconds)
                retries += 1
            except AssertionError as e:
                print(f"{spec_id}: SKIPPED -- {e}")
                skipped = True
                break
            except APIStatusError as e:
                print(f"{spec_id}: SKIPPED -- request too large for this account's TPM limit: {e}")
                skipped = True
                break

        if skipped:
            continue
        if row_result is None:
            print(f"{spec_id}: giving up after 3 retries (likely daily quota). "
                  f"Progress saved -- rerun this script to continue.")
            break

        results.append(row_result)
        pd.DataFrame(results).to_csv(OUTPUT_FILE, index=False)   # save after every row
        done_spec_ids.add(spec_id)
        time.sleep(T)  # long pause -- must fully clear the per-minute token budget between rows,
                         # since each row can issue up to 5 LLM calls with large context payloads

    df = pd.DataFrame(results)
    if len(df) == 0:
        print("\nNo rows timed yet -- nothing to summarize.")
        return
    df.to_csv(OUTPUT_FILE, index=False)

    print(f"\n=== Latency summary (n={len(df)} of {len(sample)} planned) ===\n")
    print("System A:")
    print(f"  retrieval:  mean={df['A_retrieval_ms'].mean():.0f}ms")
    print(f"  generation: mean={df['A_generation_ms'].mean():.0f}ms")
    print(f"  total:      mean={df['A_total_ms'].mean():.0f}ms")

    print("\nSystem B:")
    print(f"  retrieval:      mean={df['B_retrieval_ms'].mean():.0f}ms")
    print(f"  evaluator:      mean={df['B_evaluator_ms'].mean():.0f}ms")
    print(f"  rewrite cycle:  mean={df['B_rewrite_cycle_ms'].mean():.0f}ms (0 for direct-path rows)")
    print(f"  generation:     mean={df['B_generation_ms'].mean():.0f}ms")
    print(f"  total:          mean={df['B_total_ms'].mean():.0f}ms")

    print(f"\nOverhead: System B is {df['B_total_ms'].mean() - df['A_total_ms'].mean():.0f}ms slower than System A on average (T = {T})")
    print(f"\nPath breakdown:")
    print(df['B_path'].value_counts())
    print(f"\nSaved to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()